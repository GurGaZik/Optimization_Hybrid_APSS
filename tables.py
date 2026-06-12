from __future__ import annotations

import math

import numpy as np
import pandas as pd
from PySide6.QtCore import Qt
from PySide6.QtWidgets import (
    QAbstractItemView,
    QFileDialog,
    QHBoxLayout,
    QHeaderView,
    QMessageBox,
    QPushButton,
    QTabBar,
    QTabWidget,
    QTableWidget,
    QTableWidgetItem,
    QVBoxLayout,
    QWidget,
)
from openpyxl.styles import Alignment, Font

from data_loader import SourceData
from energy_models import SimulationResult


MONTH_NAMES = {
    1: "Январь",
    2: "Февраль",
    3: "Март",
    4: "Апрель",
    5: "Май",
    6: "Июнь",
    7: "Июль",
    8: "Август",
    9: "Сентябрь",
    10: "Октябрь",
    11: "Ноябрь",
    12: "Декабрь",
}

ALGORITHMS = ["Greedy", "GWO", "DE"]
TABLE_CHOICES = [
    ("existing_year", "Годовая исходная"),
    ("final_year", "Годовая итоговая"),
    ("jan15", "15 января"),
    ("apr15", "15 апреля"),
    ("jul15", "15 июля"),
    ("oct15", "15 октября"),
]


class TablesPanel(QWidget):

    def __init__(self) -> None:
        super().__init__()
        outer = QVBoxLayout(self)
        outer.setContentsMargins(8, 8, 8, 8)

        self.open_tables: dict[str, dict[str, pd.DataFrame]] = {
            algorithm: {} for algorithm in ALGORITHMS
        }

        export_row = QHBoxLayout()
        self.export_current_btn = QPushButton("Скачать текущую таблицу в Excel")
        self.export_current_btn.clicked.connect(self.export_current_table)
        self.export_all_btn = QPushButton("Скачать все таблицы в Excel")
        self.export_all_btn.clicked.connect(self.export_all_tables)
        export_row.addWidget(self.export_current_btn)
        export_row.addWidget(self.export_all_btn)
        export_row.addStretch(1)
        outer.addLayout(export_row)

        self.algorithm_tabs = QTabWidget()
        self.inner_tabs: dict[str, QTabWidget] = {}
        for algorithm in ALGORITHMS:
            inner = QTabWidget()
            inner.setDocumentMode(False)
            inner.setMovable(True)
            inner.setTabsClosable(False)
            inner.setStyleSheet(_tab_close_style())
            self.inner_tabs[algorithm] = inner
            self.algorithm_tabs.addTab(inner, algorithm)
        outer.addWidget(self.algorithm_tabs, 1)

    def set_message(self, message: str) -> None:
        return None

    def clear_tables(self, algorithm: str | None = None) -> None:
        targets = [algorithm] if algorithm else ALGORITHMS
        for alg in targets:
            self.open_tables.setdefault(alg, {}).clear()
            inner = self.inner_tabs.get(alg)
            if inner is None:
                continue
            while inner.count():
                widget = inner.widget(0)
                inner.removeTab(0)
                if widget is not None:
                    widget.setParent(None)
                    widget.deleteLater()

    def clear_algorithm(self, algorithm: str) -> None:
        self.clear_tables(algorithm)

    def show_algorithm(self, data: SourceData, result: SimulationResult, algorithm: str) -> None:
        self.clear_tables(algorithm)
        for key, title in TABLE_CHOICES:
            df = build_table_by_key(data, result, key)
            self.open_table(algorithm, title, df)
        self.algorithm_tabs.setCurrentWidget(self.inner_tabs[algorithm])

    def open_table(self, algorithm: str, title: str, df: pd.DataFrame) -> None:
        inner = self.inner_tabs.get(algorithm)
        if inner is None:
            raise ValueError(f"Неизвестный алгоритм для вкладки таблиц: {algorithm}")

        self.open_tables.setdefault(algorithm, {})[title] = df.copy()

        existing_index = _find_tab(inner, title)
        if existing_index >= 0:
            _close_tab(inner, existing_index)

        page = QWidget()
        layout = QVBoxLayout(page)
        layout.setContentsMargins(10, 10, 10, 10)
        layout.addWidget(_dataframe_to_table_widget(df), 1)

        index = inner.addTab(page, title)
        _install_close_button(inner, index)
        inner.setCurrentIndex(index)
        self.algorithm_tabs.setCurrentWidget(inner)

    def export_current_table(self) -> None:
        algorithm = self.algorithm_tabs.tabText(self.algorithm_tabs.currentIndex())
        inner = self.inner_tabs.get(algorithm)
        if inner is None or inner.currentIndex() < 0:
            QMessageBox.information(self, "Экспорт таблицы", "Нет открытой таблицы для экспорта.")
            return

        title = inner.tabText(inner.currentIndex())
        df = self.open_tables.get(algorithm, {}).get(title)
        if df is None:
            QMessageBox.information(self, "Экспорт таблицы", "Данные текущей таблицы не найдены.")
            return

        path, _ = QFileDialog.getSaveFileName(
            self,
            "Сохранить таблицу в Excel",
            _safe_file_name(f"{algorithm}_{title}.xlsx"),
            "Excel (*.xlsx)",
        )
        if not path:
            return

        _write_tables_to_excel(path, [(title, df)])
        QMessageBox.information(self, "Экспорт таблицы", "Таблица сохранена в Excel.")

    def export_all_tables(self) -> None:
        tables: list[tuple[str, pd.DataFrame]] = []
        for algorithm, table_map in self.open_tables.items():
            for title, df in table_map.items():
                tables.append((f"{algorithm}_{title}", df))

        if not tables:
            QMessageBox.information(self, "Экспорт таблиц", "Нет открытых таблиц для экспорта.")
            return

        path, _ = QFileDialog.getSaveFileName(
            self,
            "Сохранить все таблицы в Excel",
            "tables_export.xlsx",
            "Excel (*.xlsx)",
        )
        if not path:
            return

        _write_tables_to_excel(path, tables)
        QMessageBox.information(self, "Экспорт таблиц", "Все открытые таблицы сохранены в Excel.")


def build_table_by_key(data: SourceData, result: SimulationResult | None, key: str) -> pd.DataFrame:
    if key == "existing_year":
        return build_existing_year_table(data)
    if result is None:
        raise ValueError("Эта таблица доступна только после запуска оптимизации.")
    if key == "final_year":
        return build_final_year_table(result)
    if key == "jan15":
        return build_final_day_table(result, 1, 15)
    if key == "apr15":
        return build_final_day_table(result, 4, 15)
    if key == "jul15":
        return build_final_day_table(result, 7, 15)
    if key == "oct15":
        return build_final_day_table(result, 10, 15)
    raise ValueError("Неизвестный тип таблицы.")


def build_existing_year_table(data: SourceData) -> pd.DataFrame:
    load_df = data.load_year.copy()
    mges_df = data.mges_year.copy()
    des_df = data.des_year.copy()

    df = load_df.merge(
        mges_df[["month_number", "mges1_kwh", "mges2_kwh"]],
        on="month_number",
        how="left",
    )
    df = df.merge(des_df[["month_number", "des_kwh"]], on="month_number", how="left")
    for col in ["mges1_kwh", "mges2_kwh", "des_kwh", "load_kwh"]:
        df[col] = _numeric_series(df[col])

    out = pd.DataFrame(
        {
            "Месяц": df["month_number"].map(MONTH_NAMES),
            "Wдэу1, кВт*ч": df["des_kwh"] / 2.0,
            "Wдэу2, кВт*ч": df["des_kwh"] / 2.0,
            "Wмгэс1, кВт*ч": df["mges1_kwh"],
            "Wмгэс2, кВт*ч": df["mges2_kwh"],
            "Wнепокр, кВт*ч": 0.0,
            "Wпотр кВт*ч": df["load_kwh"],
        }
    )
    return _round_and_add_total_row(out, label_column="Месяц")


def build_final_year_table(result: SimulationResult) -> pd.DataFrame:
    df = _with_diesel_split_columns(result.hourly)
    dt = pd.to_datetime(df["datetime"])
    df = df.assign(month=dt.dt.month)

    for col in ["ballast_kwh", "unserved_kwh"]:
        if col not in df.columns:
            df[col] = 0.0

    grouped = df.groupby("month", as_index=False)[
        [
            "diesel1_kwh",
            "diesel2_kwh",
            "mges1_kwh",
            "mges2_kwh",
            "wind_used_kwh",
            "battery_discharge_kwh",
            "battery_charge_kwh",
            "ballast_kwh",
            "unserved_kwh",
            "load_kwh",
        ]
    ].sum()

    out = pd.DataFrame(
        {
            "Месяц": grouped["month"].map(MONTH_NAMES),
            "Wдэу1, кВт*ч": grouped["diesel1_kwh"],
            "Wдэу2, кВт*ч": grouped["diesel2_kwh"],
            "Wмгэс1, кВт*ч": grouped["mges1_kwh"],
            "Wмгэс2, кВт*ч": grouped["mges2_kwh"],
            "Wвэу, кВт*ч": grouped["wind_used_kwh"],
            "Wакб_раз, кВт*ч": grouped["battery_discharge_kwh"],
            "Wакб_зар, кВт*ч": -grouped["battery_charge_kwh"],
            "Wбалласт, кВт*ч": grouped["ballast_kwh"],
            "Wнепокр, кВт*ч": grouped["unserved_kwh"],
            "Wпотр кВт*ч": grouped["load_kwh"],
        }
    )
    return _round_and_add_total_row(out, label_column="Месяц")


def build_final_day_table(result: SimulationResult, month: int, day: int = 15) -> pd.DataFrame:
    df = _with_diesel_split_columns(result.hourly)
    dt = pd.to_datetime(df["datetime"])
    mask = (dt.dt.month == int(month)) & (dt.dt.day == int(day))
    day_df = df.loc[mask].copy()
    if day_df.empty:
        raise ValueError(f"В результатах нет данных за {day:02d}.{month:02d}.")

    for col in ["ballast_kwh", "unserved_kwh"]:
        if col not in day_df.columns:
            day_df[col] = 0.0

    day_df["hour"] = pd.to_datetime(day_df["datetime"]).dt.hour
    day_df = day_df.sort_values("hour")

    out = pd.DataFrame(
        {
            "Час": day_df["hour"].astype(int),
            "Pдэу1, кВт": day_df["diesel1_kwh"],
            "Pдэу2, кВт": day_df["diesel2_kwh"],
            "Pмгэс1, кВт": day_df["mges1_kwh"],
            "Pмгэс2, кВт": day_df["mges2_kwh"],
            "Pвэу, кВт": day_df["wind_used_kwh"],
            "Pакб_раз, кВт": day_df["battery_discharge_kwh"],
            "Pакб_зар, кВт": -day_df["battery_charge_kwh"],
            "Pбалласт, кВт": day_df["ballast_kwh"],
            "Pнепокр, кВт": day_df["unserved_kwh"],
            "Pпотр кВт*ч": day_df["load_kwh"],
        }
    )
    return _round_and_add_total_row(out, label_column="Час")

def _with_diesel_split_columns(hourly: pd.DataFrame) -> pd.DataFrame:
    df = hourly.copy()
    diesel_units = _numeric_series(df["diesel_units"]).astype(int)
    per_unit = _numeric_series(df["diesel_per_unit_kw"])
    df["diesel1_kwh"] = np.where(diesel_units >= 1, per_unit, 0.0)
    df["diesel2_kwh"] = np.where(diesel_units >= 2, per_unit, 0.0)
    return df


def _numeric_series(series: pd.Series) -> pd.Series:
    values = []
    for value in series.tolist():
        try:
            number = float(str(value).replace(" ", "").replace(",", "."))
            if not math.isfinite(number):
                number = 0.0
        except Exception:
            number = 0.0
        values.append(number)
    return pd.Series(values, index=series.index, dtype=float)


def _safe_number(value: object) -> float:
    try:
        number = float(str(value).replace(" ", "").replace(",", "."))
        if not math.isfinite(number):
            return 0.0
        return number
    except Exception:
        return 0.0


def _round_and_add_total_row(df: pd.DataFrame, label_column: str) -> pd.DataFrame:
    out = df.copy()
    numeric_cols = [col for col in out.columns if col != label_column]
    for col in numeric_cols:
        out[col] = [_round1(_safe_number(value)) for value in out[col].tolist()]

    total_row: dict[str, object] = {label_column: "Сумма"}
    for col in numeric_cols:
        total_row[col] = _round1(sum(_safe_number(value) for value in out[col].tolist()))
    return pd.concat([out, pd.DataFrame([total_row])], ignore_index=True)


def _round1(value: float) -> float:
    return round(float(value), 1)


def dataframe_to_table_widget(df: pd.DataFrame) -> QTableWidget:
    return _dataframe_to_table_widget(df)


def _dataframe_to_table_widget(df: pd.DataFrame) -> QTableWidget:
    table = QTableWidget()
    table.setRowCount(len(df))
    table.setColumnCount(len(df.columns))
    table.setHorizontalHeaderLabels([str(col) for col in df.columns])
    table.setAlternatingRowColors(True)
    table.setSortingEnabled(False)
    table.setEditTriggers(QAbstractItemView.NoEditTriggers)
    table.setSelectionBehavior(QAbstractItemView.SelectItems)
    table.setSelectionMode(QAbstractItemView.ExtendedSelection)

    for row_idx in range(len(df)):
        is_total = row_idx == len(df) - 1 or str(df.iloc[row_idx, 0]) == "Сумма"
        for col_idx, col_name in enumerate(df.columns):
            value = df.iloc[row_idx, col_idx]
            text = _format_cell_value(value, is_label_col=(col_idx == 0))
            item = QTableWidgetItem(str(text))
            if col_idx > 0:
                item.setTextAlignment(Qt.AlignRight | Qt.AlignVCenter)
            else:
                item.setTextAlignment(Qt.AlignLeft | Qt.AlignVCenter)
            if is_total:
                font = item.font()
                font.setBold(True)
                item.setFont(font)
            table.setItem(row_idx, col_idx, item)

    header = table.horizontalHeader()
    header.setSectionResizeMode(QHeaderView.Stretch)
    header.setStretchLastSection(False)
    table.verticalHeader().setVisible(False)
    table.resizeRowsToContents()
    return table


def _format_cell_value(value: object, is_label_col: bool) -> str:
    if is_label_col:
        return str(value)
    return _fmt_number(_safe_number(value), 1)


def _fmt_number(value: float, decimals: int = 1) -> str:
    return f"{float(value):,.{decimals}f}".replace(",", " ").replace(".", ",")


def _find_tab(tab_widget: QTabWidget, title: str) -> int:
    for i in range(tab_widget.count()):
        if tab_widget.tabText(i) == title:
            return i
    return -1


def _close_tab(tab_widget: QTabWidget, index: int) -> None:
    widget = tab_widget.widget(index)
    tab_widget.removeTab(index)
    if widget is not None:
        widget.setParent(None)
        widget.deleteLater()


def _install_close_button(tab_widget: QTabWidget, index: int) -> None:
    button = QPushButton("×")
    button.setFixedSize(24, 24)
    button.setToolTip("Закрыть вкладку")
    button.setStyleSheet(
        "QPushButton { color: black; font-size: 18px; font-weight: bold; border: none; background: transparent; }"
        "QPushButton:hover { background: #dddddd; border-radius: 11px; }"
    )
    button.clicked.connect(lambda _checked=False, tw=tab_widget, page=tab_widget.widget(index): _close_tab_by_widget(tw, page))
    position = QTabBar.ButtonPosition.RightSide
    tab_widget.tabBar().setTabButton(index, position, button)


def _close_tab_by_widget(tab_widget: QTabWidget, page: QWidget) -> None:
    index = tab_widget.indexOf(page)
    if index >= 0:
        _close_tab(tab_widget, index)


def _write_tables_to_excel(path: str, tables: list[tuple[str, pd.DataFrame]]) -> None:
    if not path.lower().endswith(".xlsx"):
        path += ".xlsx"

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        used_sheet_names: set[str] = set()

        for title, df in tables:
            sheet_name = _safe_sheet_name(title, used_sheet_names)
            df.to_excel(writer, sheet_name=sheet_name, index=False)

            worksheet = writer.sheets[sheet_name]
            worksheet.freeze_panes = "A2"

            for cell in worksheet[1]:
                cell.font = Font(bold=True)
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

            for row in worksheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(vertical="center", wrap_text=True)
                    if isinstance(cell.value, (int, float)):
                        cell.number_format = "#,##0.0"

            if worksheet.max_row >= 2:
                for cell in worksheet[worksheet.max_row]:
                    cell.font = Font(bold=True)

            for column_cells in worksheet.columns:
                max_length = 0
                column_letter = column_cells[0].column_letter
                for cell in column_cells:
                    value = "" if cell.value is None else str(cell.value)
                    max_length = max(max_length, len(value))
                worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 38)


def _safe_sheet_name(name: str, used_names: set[str]) -> str:
    for ch in r'[]:*?/\\':
        name = name.replace(ch, "_")

    name = name[:31] or "Таблица"
    base_name = name
    counter = 1

    while name in used_names:
        suffix = f"_{counter}"
        name = base_name[:31 - len(suffix)] + suffix
        counter += 1

    used_names.add(name)
    return name


def _safe_file_name(name: str) -> str:
    for ch in r'<>:"/\\|?*':
        name = name.replace(ch, "_")
    return name


def _tab_close_style() -> str:
    return """
    QTabBar::tab { padding: 7px 34px 7px 10px; }
    """
